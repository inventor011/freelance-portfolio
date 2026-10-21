"""Build a SAMPLE transaction auto-categorization workbook (Google Sheets / Excel compatible).
All transactions are synthetic demo data - not a real person's or client's statement."""
import random, datetime as dt
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.formatting.rule import FormulaRule
from openpyxl.worksheet.datavalidation import DataValidation
from openpyxl.utils import get_column_letter

random.seed(11)
RULES = [  # keyword (case-insensitive, matched anywhere in description), category, group
    # General rules first; more specific rules lower down WIN when several match.
    ("UPI", "Transfers (UPI)", "Expense"),
    ("NEFT", "Bank Transfer", "Expense"),
    ("ATM", "Cash Withdrawal", "Expense"),
    ("SWIGGY", "Food Delivery", "Expense"),
    ("ZOMATO", "Food Delivery", "Expense"),
    ("BLINKIT", "Groceries", "Expense"),
    ("BIGBASKET", "Groceries", "Expense"),
    ("DMART", "Groceries", "Expense"),
    ("AMAZON", "Shopping", "Expense"),
    ("FLIPKART", "Shopping", "Expense"),
    ("UBER", "Transport", "Expense"),
    ("OLA", "Transport", "Expense"),
    ("IRCTC", "Travel", "Expense"),
    ("INDIGO", "Travel", "Expense"),
    ("NETFLIX", "Subscriptions", "Expense"),
    ("SPOTIFY", "Subscriptions", "Expense"),
    ("YOUTUBE PREMIUM", "Subscriptions", "Expense"),
    ("JIO", "Mobile & Internet", "Expense"),
    ("AIRTEL", "Mobile & Internet", "Expense"),
    ("BESCOM", "Electricity", "Expense"),
    ("RENT", "Rent", "Expense"),
    ("SIP", "Investments", "Expense"),
    ("ZERODHA", "Investments", "Expense"),
    ("APOLLO", "Health", "Expense"),
    ("AMAZON PAY CASHBACK", "Cashback", "Income"),
    ("SALARY", "Salary", "Income"),
    ("INTEREST CREDIT", "Interest", "Income"),
    ("REFUND", "Refunds", "Income"),
    ("FREELANCE", "Freelance Income", "Income"),
]
TEMPL = [
    ("UPI/SWIGGY/{n}/Order", -1, (180, 650)), ("UPI/ZOMATO ORDER {n}", -1, (150, 700)),
    ("UPI/BLINKIT/{n}", -1, (200, 1500)), ("POS DMART BANGALORE {n}", -1, (800, 3500)),
    ("AMAZON SELLER SERVICES {n}", -1, (300, 4000)), ("FLIPKART INTERNET {n}", -1, (500, 6000)),
    ("UPI/UBER INDIA/{n}", -1, (120, 600)), ("UPI/OLA CABS/{n}", -1, (100, 500)),
    ("IRCTC E-TICKET {n}", -1, (400, 2500)), ("NETFLIX.COM SUBSCRIPTION", -1, (199, 199)),
    ("SPOTIFY INDIA", -1, (119, 119)), ("JIO PREPAID RECHARGE", -1, (299, 299)),
    ("BESCOM ELECTRICITY BILL {n}", -1, (900, 2200)), ("NEFT RENT TO LANDLORD", -1, (15000, 15000)),
    ("ACH HDFC MF SIP {n}", -1, (5000, 5000)), ("ATM WDL {n}", -1, (2000, 5000)),
    ("APOLLO PHARMACY {n}", -1, (150, 1200)), ("UPI/RAHUL K/{n}", -1, (200, 2000)),
    ("NEFT-SALARY ACME TECH PVT LTD", 1, (65000, 65000)), ("INTEREST CREDIT SB A/C", 1, (110, 160)),
    ("AMAZON PAY CASHBACK {n}", 1, (25, 150)), ("UPI REFUND FLIPKART {n}", 1, (300, 1500)),
    ("UPI/FREELANCE PAYMENT/{n}", 1, (1500, 6000)), ("CHQ DEP 00{n}", 1, (1000, 3000)),
]
rows = []
d0 = dt.date(2026, 9, 1)
for i in range(60):
    t, sign, (lo, hi) = random.choice(TEMPL)
    if i == 0: t, sign, (lo, hi) = TEMPL[18]
    amt = round(random.uniform(lo, hi), 0) if lo != hi else lo
    rows.append((d0 + dt.timedelta(days=random.randint(0, 59)), t.format(n=random.randint(1000, 9999)), sign * amt))
rows.sort()

wb = Workbook()
hdr_font = Font(bold=True, color="FFFFFF"); hdr_fill = PatternFill("solid", fgColor="1F2937")
thin = Side(style="thin", color="D1D5DB")

def header(ws, cols, widths):
    for i, (c, w) in enumerate(zip(cols, widths), 1):
        cell = ws.cell(row=1, column=i, value=c)
        cell.font = hdr_font; cell.fill = hdr_fill; cell.alignment = Alignment(horizontal="center")
        ws.column_dimensions[get_column_letter(i)].width = w
    ws.freeze_panes = "A2"

# --- Transactions
ws = wb.active; ws.title = "Transactions"
header(ws, ["Date", "Description", "Amount (₹)", "Category", "Group", "Matched keyword", "Status"], [12, 42, 14, 20, 11, 22, 14])
N_RULES = 200  # rule range is pre-sized so new rules need no formula edits
KW, CAT, GRP = f"Rules!$A$2:$A${N_RULES}", f"Rules!$B$2:$B${N_RULES}", f"Rules!$C$2:$C${N_RULES}"
N_TX = 1000   # formulas pre-filled for 1000 rows so pasted new transactions are tagged automatically
for r in range(2, N_TX + 2):
    if r - 2 < len(rows):
        d, desc, amt = rows[r - 2]
        ws.cell(row=r, column=1, value=d).number_format = "DD-MMM-YYYY"
        ws.cell(row=r, column=2, value=desc)
        ws.cell(row=r, column=3, value=amt).number_format = '#,##0;[Red]-#,##0'
    # LOOKUP(2^15, SEARCH(...)) = "last rule whose keyword appears anywhere in the description".
    # Works natively in Google Sheets AND Excel (no CSE / ARRAYFORMULA needed). Blank keywords are ignored.
    m = f'SEARCH({KW}&"",$B{r})/({KW}<>"")'
    ws.cell(row=r, column=6, value=f'=IF($B{r}="","",IFERROR(LOOKUP(2^15,{m},{KW}),""))')
    ws.cell(row=r, column=4, value=f'=IF($B{r}="","",IFERROR(LOOKUP(2^15,{m},{CAT}),"Uncategorized"))')
    ws.cell(row=r, column=5, value=f'=IF($B{r}="","",IF($C{r}>=0,"Income","Expense"))')
    ws.cell(row=r, column=7, value=f'=IF($B{r}="","",IF(D{r}="Uncategorized","⚠ Review",IF(IFERROR(LOOKUP(2^15,{m},{GRP}),"")<>E{r},"⚠ Check sign","OK")))')
ws.conditional_formatting.add(f"A2:G{N_TX+1}", FormulaRule(formula=['LEFT($G2,1)="⚠"'], fill=PatternFill("solid", fgColor="FEF3C7")))
ws.auto_filter.ref = f"A1:G{N_TX+1}"

# --- Rules
wr = wb.create_sheet("Rules")
header(wr, ["Keyword (matched anywhere, case-insensitive)", "Category", "Expected group", "Notes"], [40, 22, 16, 50])
for i, (k, c, g) in enumerate(RULES, 2):
    wr.cell(row=i, column=1, value=k); wr.cell(row=i, column=2, value=c); wr.cell(row=i, column=3, value=g)
wr["D2"] = "Rule order matters: if several keywords match, the LOWEST matching row wins."
wr["D3"] = "So keep broad keywords (UPI, NEFT) at the top and specific ones (SWIGGY, SALARY) below."
wr["D4"] = "Add a new rule = add a row. No formula changes needed (range is pre-sized to row 200)."
dv = DataValidation(type="list", formula1='"Income,Expense"', allow_blank=True); wr.add_data_validation(dv); dv.add("C2:C200")

# --- Summary
sm = wb.create_sheet("Summary")
header(sm, ["Category", "Total (₹)", "# Transactions"], [24, 16, 16])
cats = sorted(set(c for _, c, _ in RULES)) + ["Uncategorized"]
for i, c in enumerate(cats, 2):
    sm.cell(row=i, column=1, value=c)
    sm.cell(row=i, column=2, value=f'=SUMIFS(Transactions!$C:$C,Transactions!$D:$D,A{i})').number_format = '#,##0;[Red]-#,##0'
    sm.cell(row=i, column=3, value=f'=COUNTIFS(Transactions!$D:$D,A{i})')
last = len(cats) + 1
sm.cell(row=last + 2, column=1, value="Total income").font = Font(bold=True)
sm.cell(row=last + 2, column=2, value='=SUMIFS(Transactions!$C:$C,Transactions!$C:$C,">0")').number_format = '#,##0'
sm.cell(row=last + 3, column=1, value="Total expense").font = Font(bold=True)
sm.cell(row=last + 3, column=2, value='=SUMIFS(Transactions!$C:$C,Transactions!$C:$C,"<0")').number_format = '#,##0;[Red]-#,##0'
sm.cell(row=last + 4, column=1, value="Rows needing review").font = Font(bold=True)
sm.cell(row=last + 4, column=2, value='=COUNTIF(Transactions!$G:$G,"⚠*")')

ab = wb.create_sheet("About")
ab.column_dimensions["A"].width = 110
for i, line in enumerate([
    "SAMPLE / DEMO WORKBOOK - synthetic transactions, not real client data.",
    "Built by Ayush Kumar as a portfolio sample of rule-based transaction categorization.",
    "Transactions: paste raw rows into columns A-C; columns D-G tag themselves.",
    "Rules: keyword -> category table. Edit freely; formulas read the whole range.",
    "Summary: totals per category, income vs expense, and a count of rows flagged for review.",
], 1):
    ab.cell(row=i, column=1, value=line)
wb.move_sheet("About", offset=-3)
wb.save("transaction_categorizer_SAMPLE.xlsx")
print("rows", len(rows))
